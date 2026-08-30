"""
Open the hand-written .xlsx with a real reader and check it satisfies the analyzer's
contract: one sheet named after the product, the seven header strings matching exactly,
and numeric cells that came through as numbers rather than text.
"""
import json, subprocess, sys, os, tempfile

HEADER = ['月別推移', '売上金額', '購入個数', '購入件数', '総顧客数', '新規顧客数', '既存顧客数']

repo = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
tmp = os.path.join(tempfile.gettempdir(), 'basket_demo_export_check.xlsx')

proc = subprocess.run(
    ['node', os.path.join(repo, 'tools', 'check_xlsx_export.js'), tmp],
    capture_output=True, text=True, encoding='utf-8'
)
if proc.returncode != 0:
    print(proc.stdout)
    print(proc.stderr, file=sys.stderr)
    sys.exit('building the export failed')

meta = json.loads(proc.stdout.strip().splitlines()[-1])

try:
    import openpyxl
except ImportError:
    sys.exit('openpyxl is required: pip install openpyxl')

wb = openpyxl.load_workbook(tmp)
problems = []

if wb.sheetnames != [meta['sheetName']]:
    problems.append('sheet names %r, expected [%r]' % (wb.sheetnames, meta['sheetName']))

ws = wb[wb.sheetnames[0]]
rows = list(ws.iter_rows(values_only=True))

if not rows:
    problems.append('the sheet is empty')
else:
    if list(rows[0]) != HEADER:
        problems.append('header row is %r' % (list(rows[0]),))
    if len(rows) != meta['rows']:
        problems.append('%d rows, expected %d' % (len(rows), meta['rows']))
    for r, row in enumerate(rows[1:], start=2):
        if not isinstance(row[0], str):
            problems.append('row %d: month label is %r, expected text' % (r, row[0]))
            break
        for c, value in enumerate(row[1:], start=2):
            if not isinstance(value, int):
                problems.append('row %d col %d: %r is not a number' % (r, c, value))
                break
        else:
            continue
        break

first = list(rows[1]) if len(rows) > 1 else []
expected_first = meta['firstDataRow']
if first and first != expected_first:
    problems.append('first data row %r != %r' % (first, expected_first))

wb.close()
os.remove(tmp)

print('sheet name : %s' % meta['sheetName'])
print('rows       : %d (1 header + %d months)' % (meta['rows'], meta['rows'] - 1))
print('size       : %d bytes' % meta['bytes'])
print('first row  : %s' % (first,))

if problems:
    print('\nFAILURES (%d):' % len(problems))
    for p in problems:
        print('  ' + p)
    sys.exit(1)
print('\nthe exported workbook satisfies the analyzer contract')
